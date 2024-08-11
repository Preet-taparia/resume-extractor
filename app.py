import os
import re
from flask import Flask, render_template, request, send_file
from werkzeug.utils import secure_filename
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import Workbook

app = Flask(__name__)

UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'pdf', 'doc', 'docx'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_information(file_path, file_extension):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    phone_pattern = r'(\+?\d{2}[-\s]?\d{10}|\(\+?\d{2}\)\s?\d{10}|\+?\d{2}\s?\d{4}[-\s]?\d{4}|\d{4}[-\s]?\d{3}[-\s]?\d{3})'

    text = ""
    emails = []
    phones = []

    try:
        if file_extension == "pdf":
            with open(file_path, 'rb') as f:
                reader = PdfReader(f)
                for page in reader.pages:
                    text += page.extract_text()

        elif file_extension in ["doc", "docx"]:
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text

        else:
            return ("", "", "")

        emails = re.findall(email_pattern, text)
        phones = re.findall(phone_pattern, text)

        return emails, phones, text

    except Exception as e:
        print(f"Error extracting information from {file_path}: {e}")
        return ("", "", "")


@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return 'No file part'
    
    files = request.files.getlist('file')
    
    if len(files) == 0:
        return 'No files selected'
    
    wb = Workbook()
    ws = wb.active
    ws.append(['Email', 'Phone', 'Text'])
    
    for file in files:
        if allowed_file(file.filename):
            filename = secure_filename(file.filename)
            file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(file_path)
            
            emails, phones, text = extract_information(file_path, filename.rsplit('.', 1)[1].lower())
            for email, phone in zip(emails, phones):
                ws.append([email, phone, text])
            
    excel_filename = 'all_data.xlsx'
    excel_file_path = os.path.join(app.config['UPLOAD_FOLDER'], excel_filename)
    wb.save(excel_file_path)
    excel_response = send_file(excel_file_path, as_attachment=True)

    for file in os.listdir(app.config['UPLOAD_FOLDER']):
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], file)
        if os.path.isfile(file_path):
            os.remove(file_path)

    return excel_response

if __name__ == '__main__':
    app.run(debug=True)
